# -*- coding: utf-8 -*-
"""MediKiosk SIH26047 pitch deck generator - image/diagram heavy, minimal text."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------- palette ----------
TEAL      = RGBColor(0x0E, 0x7C, 0x7B)
TEAL_DARK = RGBColor(0x0B, 0x5A, 0x59)
NAVY      = RGBColor(0x14, 0x28, 0x3A)
BG        = RGBColor(0xF6, 0xFB, 0xFA)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE    = RGBColor(0xF3, 0x9C, 0x12)
RED       = RGBColor(0xE1, 0x4D, 0x43)
GREEN     = RGBColor(0x2E, 0xA8, 0x4F)
PURPLE    = RGBColor(0x6A, 0x4C, 0x93)
BLUE      = RGBColor(0x24, 0x71, 0xB8)
GRAY      = RGBColor(0x5A, 0x6B, 0x70)

SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background()
    r.shadow.inherit = False
    return s

def box(s, x, y, w, h, fill=WHITE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1.5)
    sp.shadow.inherit = False
    try:
        sp.adjustments[0] = radius
    except Exception:
        pass
    return sp

def text(s, x, y, w, h, txt, size=18, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         font="Segoe UI", anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = txt.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln
        p.alignment = align
        for run in p.runs:
            run.font.size = Pt(size); run.font.bold = bold
            run.font.color.rgb = color; run.font.name = font
    return tb

def emoji(s, x, y, w, h, ch, size=40):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.text = ch; p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.size = Pt(size); run.font.name = "Segoe UI Emoji"
        run.font.color.rgb = NAVY
    return tb

def header(s, kicker, title):
    bar = box(s, 0, 0, SW, Inches(0.16), fill=TEAL, shape=MSO_SHAPE.RECTANGLE)
    text(s, Inches(0.55), Inches(0.35), Inches(9.5), Inches(0.35),
         kicker.upper(), size=13, color=TEAL, bold=True)
    text(s, Inches(0.55), Inches(0.62), Inches(11.5), Inches(0.75),
         title, size=32, color=NAVY, bold=True)

def chip(s, x, y, w, h, label, fill=TEAL, size=13, color=WHITE):
    c = box(s, x, y, w, h, fill=fill, radius=0.5)
    tf = c.text_frame; tf.word_wrap = True
    tf.margin_left = Emu(36000); tf.margin_right = Emu(36000)
    tf.margin_top = Emu(18000); tf.margin_bottom = Emu(18000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.text = label; p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.size = Pt(size); run.font.bold = True
        run.font.color.rgb = color; run.font.name = "Segoe UI"
    return c

def arrow_right(s, x, y, size=26, color=GRAY):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, Inches(0.45), Inches(0.3))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()
    a.shadow.inherit = False
    return a

# =====================================================================
# SLIDE 1 - TITLE
# =====================================================================
s = slide(bg=NAVY)
box(s, 0, 0, SW, SH, fill=None, line=TEAL, shape=MSO_SHAPE.RECTANGLE)
# decorative circles
for cx, cy, d, col in [(11.2, -1.2, 4.2, TEAL_DARK), (-1.4, 5.2, 3.6, TEAL_DARK)]:
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = col; c.line.fill.background(); c.shadow.inherit = False

emoji(s, Inches(5.92), Inches(0.75), Inches(1.5), Inches(1.3), "\U0001F468\u200D\u2695\uFE0F", 72)
text(s, Inches(1.17), Inches(2.15), Inches(11), Inches(1.1),
     "MediKiosk", size=60, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(s, Inches(1.17), Inches(3.25), Inches(11), Inches(0.55),
     "AI-Powered Clinical History Kiosk with Live Talking Doctor",
     size=22, color=RGBColor(0xBF,0xE8,0xE5), align=PP_ALIGN.CENTER)
chip(s, Inches(4.42), Inches(4.15), Inches(4.5), Inches(0.55),
     "PS SIH26047  \u2022  Ministry of AYUSH", fill=TEAL, size=15)
chip(s, Inches(4.87), Inches(4.85), Inches(3.6), Inches(0.5),
     "Smart India Hackathon 2026", fill=RGBColor(0x1F,0x3A,0x52), size=13)
row_y = Inches(5.95); labels = ["\U0001F3A4 Voice+Touch", "\U0001F9FE Scan Reports",
                                "\U0001FA7A FHIR/HIS", "\U0001F6AA OPD Ready"]
w_chip, gap = Inches(2.55), Inches(0.25)
total = w_chip*4 + gap*3
x = (SW - total)//2
for lb in labels:
    chip(s, x, row_y, w_chip, Inches(0.5), lb, fill=RGBColor(0x1F,0x3A,0x52), size=13)
    x += w_chip + gap

# =====================================================================
# SLIDE 2 - PROBLEM (icon grid, no paragraphs)
# =====================================================================
s = slide()
header(s, "The Problem Today", "OPD Case-Taking Is Broken")
cards = [
    ("\U0001F46B", "Long queues,\n~5 min per patient", ORANGE),
    ("\u270D\uFE0F", "Doctor writes\nhistory by hand", BLUE),
    ("\U0001F5C3\uFE0F", "Old prescriptions\n& reports are paper", PURPLE),
    ("\U0001F9E0", "Patient forgets meds,\ndoses, allergies", RED),
    ("\U0001F5E3\uFE0F", "Language & literacy\nbarriers", GREEN),
    ("\u23F0", "Red-flag symptoms\nwait in routine queue", TEAL),
]
cw, ch_, gx, gy = Inches(3.85), Inches(2.05), Inches(0.35), Inches(0.35)
x0, y0 = (SW - (cw*3+gx*2))//2, Inches(1.75)
for i, (ic, cap, col) in enumerate(cards):
    r, cidx = divmod(i, 3)
    x, y = x0 + cidx*(cw+gx), y0 + r*(ch_+gy)
    card = box(s, x, y, cw, ch_, fill=WHITE)
    strip = box(s, x, y, Inches(0.14), ch_, fill=col, radius=0.5)
    emoji(s, x+Inches(0.35), y+Inches(0.45), Inches(1.1), Inches(1.1), ic, 44)
    text(s, x+Inches(1.55), y+Inches(0.35), cw-Inches(1.8), ch_-Inches(0.6),
         cap, size=17, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(0.55), Inches(6.65), Inches(12.2), Inches(0.5),
     "\u23F1  Doctors spend consultation time collecting history instead of treating",
     size=16, color=RED, bold=True, align=PP_ALIGN.CENTER)

# =====================================================================
# SLIDE 3 - SOLUTION: PATIENT JOURNEY (visual flow)
# =====================================================================
s = slide()
header(s, "Proposed Solution", "One Kiosk Walk-In \u2014 5 Steps")
steps = [
    ("1", "\U0001F194", "Identify", "ABHA ID /\nnew registration", TEAL),
    ("2", "\U0001F506", "Consent", "Audio-guided,\nEN / HI", GREEN),
    ("3", "\U0001F9D1\u200D\u2695\uFE0F", "Talk", "Dr. Sahayak\nAI interview", BLUE),
    ("4", "\U0001F4F8", "Scan", "Prescriptions &\nlab reports", PURPLE),
    ("5", "\U0001F4CB", "Summary", "Structured draft\nto physician", ORANGE),
]
cw, gx = Inches(2.15), Inches(0.32)
total = cw*5 + gx*4
x = (SW - total)//2; y = Inches(2.1)
for i, (n, ic, t1, t2, col) in enumerate(steps):
    card = box(s, x, y, cw, Inches(3.0), fill=WHITE)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, x+cw//2-Inches(0.3), y-Inches(0.3), Inches(0.6), Inches(0.6))
    circ.fill.solid(); circ.fill.fore_color.rgb = col; circ.line.color.rgb = WHITE
    circ.line.width = Pt(2.5); circ.shadow.inherit = False
    ct = circ.text_frame; ct.paragraphs[0].text = n; ct.paragraphs[0].alignment = PP_ALIGN.CENTER
    for rn in ct.paragraphs[0].runs:
        rn.font.bold = True; rn.font.size = Pt(20); rn.font.color.rgb = WHITE
    emoji(s, x, y+Inches(0.45), cw, Inches(1.0), ic, 48)
    text(s, x, y+Inches(1.5), cw, Inches(0.45), t1, size=20, bold=True,
         color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x+Inches(0.1), y+Inches(2.0), cw-Inches(0.2), Inches(0.9), t2,
         size=13, color=GRAY, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow_right(s, x+cw+Inches(0.02), y+Inches(1.35))
    x += cw + gx
band = box(s, Inches(1.2), Inches(5.75), Inches(10.9), Inches(1.0), fill=TEAL)
text(s, Inches(1.5), Inches(5.95), Inches(10.3), Inches(0.6),
     "\u2705  Patient walks out \u2022 Doctor gets ready-made structured history \u2022 Session data auto-purged",
     size=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# =====================================================================
# SLIDE 4 - THE TALKING DOCTOR (Dr. Sahayak)
# =====================================================================
s = slide()
header(s, "The Differentiator", "Meet Dr. Sahayak \u2014 The Talking AI Doctor")
# left: avatar mockup
av = box(s, Inches(0.8), Inches(1.8), Inches(4.6), Inches(4.9), fill=NAVY, radius=0.06)
screen = box(s, Inches(1.1), Inches(2.05), Inches(4.0), Inches(3.3), fill=RGBColor(0x1F,0x3A,0x52), radius=0.04)
emoji(s, Inches(2.3), Inches(2.6), Inches(1.6), Inches(1.6), "\U0001F9D1\u200D\u2695\uFE0F", 80)
wave = box(s, Inches(1.6), Inches(4.55), Inches(3.0), Inches(0.5), fill=TEAL, radius=0.5)
text(s, Inches(1.6), Inches(4.62), Inches(3.0), Inches(0.4), "\U0001F5E3\uFE0F  Lip-synced video",
     size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
mic = box(s, Inches(2.55), Inches(5.55), Inches(1.1), Inches(1.1), fill=RED, shape=MSO_SHAPE.OVAL)
mt = mic.text_frame; mt.paragraphs[0].text = "\U0001F3A4"; mt.paragraphs[0].alignment = PP_ALIGN.CENTER
for rn in mt.paragraphs[0].runs: rn.font.size = Pt(34)
text(s, Inches(1.1), Inches(5.5), Inches(4.0), Inches(0.4), "", size=10)
# right: two brains
bx = Inches(6.0); bw = Inches(6.6)
c1 = box(s, bx, Inches(1.8), bw, Inches(2.25), fill=WHITE)
strip = box(s, bx, Inches(1.8), Inches(0.14), Inches(2.25), fill=BLUE, radius=0.5)
emoji(s, bx+Inches(0.3), Inches(2.1), Inches(1.0), Inches(1.0), "\U0001F39E\uFE0F", 44)
text(s, bx+Inches(1.5), Inches(1.98), bw-Inches(1.8), Inches(0.45),
     "Video Mode (primary)", size=20, bold=True, color=BLUE)
text(s, bx+Inches(1.5), Inches(2.5), bw-Inches(1.8), Inches(1.4),
     "Photoreal avatar conducts SOCRATES-style\nadaptive questioning \u2014 voice in, face out",
     size=15, color=GRAY)
c2 = box(s, bx, Inches(4.45), bw, Inches(2.25), fill=WHITE)
strip = box(s, bx, Inches(4.45), Inches(0.14), Inches(2.25), fill=PURPLE, radius=0.5)
emoji(s, bx+Inches(0.3), Inches(4.75), Inches(1.0), Inches(1.0), "\u261F\uFE0F", 44)
text(s, bx+Inches(1.5), Inches(4.63), bw-Inches(1.8), Inches(0.45),
     "Touch Mode (fallback)", size=20, bold=True, color=PURPLE)
text(s, bx+Inches(1.5), Inches(5.15), bw-Inches(1.8), Inches(1.4),
     "Icon MCQ taps + voice \u2014 works when noise,\nlow literacy or avatar failure hits",
     size=15, color=GRAY)
chip(s, bx+Inches(0.3), Inches(6.85), Inches(6.0), Inches(0.45),
     "\U0001F54D  AYUSH mode: full Dashavidha Pariksha capture", fill=GREEN, size=13)

# =====================================================================
# SLIDE 5 - ARCHITECTURE DIAGRAM
# =====================================================================
s = slide()
header(s, "Technical Approach", "Architecture")
# kiosk frontend layer
fe = box(s, Inches(0.7), Inches(1.75), Inches(11.9), Inches(1.35), fill=WHITE, line=TEAL)
text(s, Inches(0.95), Inches(1.83), Inches(4), Inches(0.35), "KIOSK FRONTEND (React)", size=14, bold=True, color=TEAL)
fe_icons = [("\U0001F9D1\u200D\u2695\uFE0F","Avatar"),("\U0001F3A4","Voice"),("\u261F\uFE0F","Touch"),("\U0001F4F8","Doc Scan")]
fx = Inches(0.95)
for ic, lb in fe_icons:
    emoji(s, fx, Inches(2.25), Inches(0.7), Inches(0.7), ic, 30)
    text(s, fx+Inches(0.68), Inches(2.38), Inches(1.2), Inches(0.4), lb, size=13, bold=True, color=NAVY)
    fx += Inches(1.85)
arrow_down = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, SW//2-Inches(0.2), Inches(3.15), Inches(0.4), Inches(0.5))
arrow_down.fill.solid(); arrow_down.fill.fore_color.rgb = GRAY; arrow_down.line.fill.background(); arrow_down.shadow.inherit=False
# backend modules
mods = [
    ("M1 \u2022 Interview Engine", "Ontology branching\nRed-flag detector", TEAL),
    ("M2 \u2022 Document AI", "VLM OCR extract\nTimeline + flags", BLUE),
    ("M3 \u2022 Summary Generator", "Physician format\nDraft-only", PURPLE),
    ("M4 \u2022 Consent & ABDM", "ABHA auth\nFHIR push + wipe", GREEN),
]
bw2, gx2 = Inches(2.85), Inches(0.17)
x = Inches(0.7); y = Inches(3.75)
for t1, t2, col in mods:
    c = box(s, x, y, bw2, Inches(1.55), fill=WHITE)
    top = box(s, x, y, bw2, Inches(0.5), fill=col, radius=0.12)
    tt = top.text_frame; tt.paragraphs[0].text = t1; tt.paragraphs[0].alignment = PP_ALIGN.CENTER
    for rn in tt.paragraphs[0].runs: rn.font.bold=True; rn.font.size=Pt(13); rn.font.color.rgb=WHITE
    text(s, x+Inches(0.15), y+Inches(0.6), bw2-Inches(0.3), Inches(0.9), t2, size=13, color=GRAY, align=PP_ALIGN.CENTER)
    x += bw2 + gx2
# services row
svcs = [("\u2728 Gemini LLM/VLM"), ("\U0001F399 ASR / TTS"), ("\U0001F5C4\uFE0F Postgres"), ("\U0001F3E5 HIS (FHIR R4)")]
x = Inches(0.7)
for i,lb in enumerate(svcs):
    chip(s, x, Inches(5.75), Inches(2.85), Inches(0.6), lb, fill=NAVY, size=14)
    if i < 3:
        ln = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x+Inches(2.86), Inches(5.9), Inches(0.16), Inches(0.3))
        ln.fill.solid(); ln.fill.fore_color.rgb = GRAY; ln.line.fill.background(); ln.shadow.inherit=False
    x += bw2 + gx2
conn = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, SW//2-Inches(0.2), Inches(5.32), Inches(0.4), Inches(0.38))
conn.fill.solid(); conn.fill.fore_color.rgb = GRAY; conn.line.fill.background(); conn.shadow.inherit=False
text(s, Inches(0.7), Inches(6.6), Inches(12), Inches(0.45),
     "FastAPI backend  \u2022  Avatar is an accessibility layer \u2014 removable without changing the core engine",
     size=14, color=GRAY, align=PP_ALIGN.CENTER)

# =====================================================================
# SLIDE 6 - DOC SCAN -> TIMELINE (feature visual)
# =====================================================================
s = slide()
header(s, "Document Intelligence", "Paper Prescriptions \u2192 Structured Timeline")
# crumpled paper -> scan -> structured json -> timeline
flow = [
    ("\U0001F4C4", "Scan / upload\nany prescription", ORANGE),
    ("\U0001F52D", "Gemini Vision\nOCR extraction", BLUE),
    ("\u2699\uFE0F", "Meds \u2022 Labs \u2022 Dx\nstructured JSON", PURPLE),
    ("\U0001F4C5", "Chronological\nmedical timeline", GREEN),
]
cw, gx = Inches(2.6), Inches(0.55)
x = Inches(0.75); y = Inches(2.0)
for ic, t2, col in flow:
    c = box(s, x, y, cw, Inches(2.3), fill=WHITE)
    ring = s.shapes.add_shape(MSO_SHAPE.OVAL, x+cw//2-Inches(0.65), y+Inches(0.25), Inches(1.3), Inches(1.3))
    ring.fill.solid(); ring.fill.fore_color.rgb = col; ring.line.fill.background(); ring.shadow.inherit=False
    et = ring.text_frame; et.paragraphs[0].text = ic; et.paragraphs[0].alignment=PP_ALIGN.CENTER
    for rn in et.paragraphs[0].runs: rn.font.size=Pt(36)
    text(s, x+Inches(0.15), y+Inches(1.65), cw-Inches(0.3), Inches(0.6), t2,
         size=13, bold=False, color=GRAY, align=PP_ALIGN.CENTER)
    x += cw + gx
    if x < Inches(12):
        pass
x = Inches(0.75)
for i in range(3):
    arrow_right(s, Inches(0.75)+(i+1)*(cw+gx)-gx+Inches(0.05), y+Inches(0.95), color=GRAY)
# flags row
fy = Inches(4.85); fw = Inches(3.7); fgx = Inches(0.45)
fx = (SW-(fw*3+fgx*2))//2
flags = [("\u26A0\uFE0F","Abnormal lab values\nhighlighted vs range", RED),
         ("\u26D4","Drug-interaction\nflagging", ORANGE),
         ("\u2753","Low confidence \u2192 asks\npatient to confirm", BLUE)]
for ic, t2, col in flags:
    c = box(s, fx, fy, fw, Inches(1.55), fill=WHITE)
    strip = box(s, fx, fy, Inches(0.14), Inches(1.55), fill=col, radius=0.5)
    emoji(s, fx+Inches(0.3), fy+Inches(0.35), Inches(0.85), Inches(0.85), ic, 34)
    text(s, fx+Inches(1.25), fy+Inches(0.25), fw-Inches(1.5), Inches(1.1),
         t2, size=14, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    fx += fw + fgx
text(s, Inches(0.55), Inches(6.65), Inches(12.2), Inches(0.45),
     "Works on handwritten \u2022 printed \u2022 Hindi-English mixed documents",
     size=15, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

# =====================================================================
# SLIDE 7 - SAFETY: RED FLAGS + PHYSICIAN IN LOOP
# =====================================================================
s = slide()
header(s, "Clinical Safety", "Never Misses an Emergency")
# alert path big
alert = box(s, Inches(0.8), Inches(1.85), Inches(6.1), Inches(4.7), fill=RGBColor(0xFD,0xED,0xEC))
emoji(s, Inches(3.0), Inches(2.1), Inches(1.7), Inches(1.3), "\U0001F6A8", 56)
text(s, Inches(1.0), Inches(3.4), Inches(5.7), Inches(0.5),
     "Radiating chest pain \u2022 Stroke FAST signs", size=17, bold=True, color=RED, align=PP_ALIGN.CENTER)
a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(3.55), Inches(3.95), Inches(0.4), Inches(0.55))
a.fill.solid(); a.fill.fore_color.rgb = RED; a.line.fill.background(); a.shadow.inherit=False
triage = box(s, Inches(1.5), Inches(4.6), Inches(4.7), Inches(0.75), fill=RED, radius=0.2)
tt = triage.text_frame; tt.paragraphs[0].text = "INSTANT PRIORITY TRIAGE ALERT"; tt.paragraphs[0].alignment=PP_ALIGN.CENTER
for rn in tt.paragraphs[0].runs: rn.font.bold=True; rn.font.size=Pt(16); rn.font.color.rgb=WHITE
text(s, Inches(1.0), Inches(5.5), Inches(5.7), Inches(0.9),
     "Interview pauses \u2192 patient jumps the queue\nRule-based \u2014 never guessed by an LLM",
     size=14, color=GRAY, align=PP_ALIGN.CENTER)
# right column: physician in loop
rx = Inches(7.3); rw = Inches(5.3)
doc = box(s, rx, Inches(1.85), rw, Inches(4.7), fill=WHITE, line=TEAL)
text(s, rx+Inches(0.3), Inches(2.0), rw-Inches(0.6), Inches(0.5),
     "\U0001F468\u200D\u2695\uFE0F  Physician-in-the-loop", size=19, bold=True, color=TEAL_DARK)
rows = [("\U0001F4C4","CC \u2192 HPI \u2192 PMH \u2192 Drugs\n\u2192 Family \u2192 ROS order"),
        ("\u270F\uFE0F","Every field editable\nAccept / edit / reject"),
        ("\U0001F4C4","One-page PDF export"),
        ("\u26A0\uFE0F","Draft only \u2014 NEVER an\nautonomous diagnosis")]
ry = Inches(2.7)
for ic, t2 in rows:
    emoji(s, rx+Inches(0.3), ry, Inches(0.7), Inches(0.7), ic, 28)
    text(s, rx+Inches(1.1), ry+Inches(0.02), rw-Inches(1.4), Inches(0.85),
         t2, size=14, color=NAVY)
    ry += Inches(0.95)

# =====================================================================
# SLIDE 8 - PRIVACY / CONSENT / ABDM
# =====================================================================
s = slide()
header(s, "Privacy by Design", "DPDP 2023 + ABDM Compliant Flow")
steps8 = [
    ("\U0001F9FE", "Identity", "ABHA / Aadhaar /\nregister new", BLUE),
    ("\U0001F50A", "Audio consent", "For every literacy\nlevel, EN / HI", GREEN),
    ("\u2611\uFE0F", "Granular toggles", "History \u2022 docs \u2022 share\nRevocable anytime", ORANGE),
    ("\U0001F3E5", "FHIR R4 push", "To hospital HIS\nlinked to ABHA", PURPLE),
    ("\U0001F5D1\uFE0F", "Auto purge", "Session wiped after\nsubmission", RED),
]
cw, gx = Inches(2.2), Inches(0.28)
total = cw*5 + gx*4
x = (SW-total)//2; y = Inches(2.2)
for ic, t1, t2, col in steps8:
    c = box(s, x, y, cw, Inches(2.7), fill=WHITE)
    dot = box(s, x, y, cw, Inches(0.14), fill=col, shape=MSO_SHAPE.RECTANGLE)
    emoji(s, x, y+Inches(0.35), cw, Inches(0.9), ic, 44)
    text(s, x, y+Inches(1.3), cw, Inches(0.45), t1, size=17, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x+Inches(0.08), y+Inches(1.78), cw-Inches(0.16), Inches(0.85), t2, size=12, color=GRAY, align=PP_ALIGN.CENTER)
    if x < (SW-total)//2 + (cw+gx)*4 - Inches(0.01):
        arrow_right(s, x+cw-Inches(0.02), y+Inches(1.2), color=GRAY)
    x += cw + gx
band = box(s, Inches(1.7), Inches(5.5), Inches(9.9), Inches(1.15), fill=TEAL)
text(s, Inches(2.0), Inches(5.66), Inches(9.3), Inches(0.85),
     "\U0001F512  Consent before processing \u2022 Data minimisation \u2022 No retention\n     Physician-in-the-loop by design \u2014 never autonomous diagnosis",
     size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# =====================================================================
# SLIDE 9 - FEASIBILITY & VIABILITY
# =====================================================================
s = slide()
header(s, "Feasibility & Viability", "Built With Proven, Affordable Tech")
left = [("React + FastAPI", "\u2705 Working prototype exists TODAY"),
        ("Gemini API", "\u2705 Free-tier friendly, Indian data residency option"),
        ("BeyondPresence avatar", "\u2705 Streaming API, graceful TTS-only fallback"),
        ("Postgres + Docker", "\u2705 Deployable on hospital LAN / cloud")]
ly = Inches(1.85)
for t1, t2 in left:
    c = box(s, Inches(0.7), ly, Inches(6.0), Inches(1.05), fill=WHITE)
    strip = box(s, Inches(0.7), ly, Inches(0.12), Inches(1.05), fill=GREEN, radius=0.5)
    text(s, Inches(1.0), ly+Inches(0.08), Inches(5.5), Inches(0.4), t1, size=16, bold=True, color=NAVY)
    text(s, Inches(1.0), ly+Inches(0.5), Inches(5.5), Inches(0.45), t2, size=13, color=GRAY)
    ly += Inches(1.2)
right = box(s, Inches(7.1), Inches(1.85), Inches(5.5), Inches(4.65), fill=NAVY)
text(s, Inches(7.4), Inches(2.05), Inches(4.9), Inches(0.5), "\U0001F4B0  Cost per kiosk", size=18, bold=True, color=WHITE)
costs = [("\U0001F4BB", "Touch screen PC", "~\u20B930k one-time"),
         ("\U0001F3A4", "Mic + speakers", "~\u20B92k"),
         ("\u2601\uFE0F", "API costs", "< \u20B91 per consult"),
         ("\U0001F69B", "Deployment", "Existing OPD desk space")]
cy = Inches(2.7)
for ic, t1, t2 in costs:
    emoji(s, Inches(7.4), cy, Inches(0.7), Inches(0.7), ic, 26)
    text(s, Inches(8.2), cy+Inches(0.02), Inches(2.2), Inches(0.4), t1, size=14, bold=True, color=WHITE)
    text(s, Inches(8.2), cy+Inches(0.4), Inches(4.2), Inches(0.4), t2, size=13, color=RGBColor(0xBF,0xE8,0xE5))
    cy += Inches(0.95)
text(s, Inches(0.7), Inches(6.7), Inches(12), Inches(0.4),
     "No new hardware invention needed \u2014 runs on commodity kiosk hardware available today",
     size=14, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

# =====================================================================
# SLIDE 10 - IMPACT & BENEFITS
# =====================================================================
s = slide()
header(s, "Impact & Benefits", "Who Wins, and How")
imps = [
    ("\U0001F468\u200D\u2695\uFE0F", "Doctors", "Full history before\npatient enters room", TEAL),
    ("\U0001F46B", "Patients", "Own language,\nzero-literacy usable", BLUE),
    ("\U0001F3DB\uFE0F", "Hospitals", "Digital records without\nchanging HIS (FHIR)", PURPLE),
    ("\U0001F52F", "AYUSH", "Dashavidha Pariksha\ndigitised at scale", GREEN),
    ("\U0001F6A8", "Emergencies", "Red-flag cases never\nsit in routine queue", RED),
    ("\U0001F4CA", "System", "Structured data for\npublic health insights", ORANGE),
]
cw, ch_, gx, gy = Inches(3.85), Inches(2.05), Inches(0.35), Inches(0.35)
x0, y0 = (SW-(cw*3+gx*2))//2, Inches(1.8)
for i, (ic, t1, t2, col) in enumerate(imps):
    r, cidx = divmod(i, 3)
    x, y = x0 + cidx*(cw+gx), y0 + r*(ch_+gy)
    c = box(s, x, y, cw, ch_, fill=WHITE)
    strip = box(s, x, y, Inches(0.14), ch_, fill=col, radius=0.5)
    emoji(s, x+Inches(0.3), y+Inches(0.5), Inches(1.0), Inches(1.0), ic, 42)
    text(s, x+Inches(1.4), y+Inches(0.25), cw-Inches(1.65), Inches(0.45), t1, size=18, bold=True, color=NAVY)
    text(s, x+Inches(1.4), y+Inches(0.75), cw-Inches(1.65), Inches(1.1), t2, size=13, color=GRAY)
text(s, Inches(0.55), Inches(6.55), Inches(12.2), Inches(0.5),
     "\U0001F9EC Evaluable: ASR accuracy \u2022 OCR field precision/recall \u2022 red-flag recall \u2022 latency < 2.5 s",
     size=15, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

# =====================================================================
# SLIDE 11 - RESEARCH & REFERENCES
# =====================================================================
s = slide()
header(s, "Research & References", "Grounded In Standards, Not Guesswork")
refs = [
    ("\U0001F4C4", "SIH Problem Statement SIH26047 \u2014 Ministry of Ayush, Patient Case-Taking Software"),
    ("\U0001F3E5", "HL7 FHIR R4 \u2014 Patient, Observation & DocumentReference resources"),
    ("\U0001F6E1\uFE0F", "DPDP Act 2023 + ABDM consent framework & Sandbox documentation"),
    ("\U0001F9EA", "SOCRATES symptom framework \u2022 Ayurveda Dashavidha Pariksha literature"),
    ("\U0001F310", "Bhashini / AI4Bharat Indic ASR-TTS APIs (named in PS)"),
    ("\U0001F4DA", "Gemini VLM document-extraction benchmarks on handwritten prescriptions"),
]
ry = Inches(1.95)
for ic, t1 in refs:
    c = box(s, Inches(0.8), ry, Inches(11.7), Inches(0.72), fill=WHITE)
    emoji(s, Inches(1.05), ry+Inches(0.08), Inches(0.6), Inches(0.6), ic, 22)
    text(s, Inches(1.8), ry+Inches(0.12), Inches(10.5), Inches(0.5), t1, size=14, color=NAVY)
    ry += Inches(0.82)
band = box(s, Inches(2.6), Inches(6.85), Inches(8.1), Inches(0.55), fill=TEAL, radius=0.5)
tb2 = band.text_frame; tb2.paragraphs[0].text = "\U0001F680 MediKiosk \u2014 History-taking in seconds, care in minutes"
tb2.paragraphs[0].alignment = PP_ALIGN.CENTER
for rn in tb2.paragraphs[0].runs: rn.font.size = Pt(15); rn.font.bold = True; rn.font.color.rgb = WHITE

out = r"C:\Users\NETHAN\Projects\MediKiosk\ppt\MediKiosk_SIH26047.pptx"
prs.save(out)
print("Saved:", out)
